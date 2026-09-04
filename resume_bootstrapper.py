"""
Resume Bootstrapper for Stage 1 (Master Profile Intake).
Extracts text from PDF, DOCX, or TXT resume files and uses structured LLM extraction
to build a fully-formed, validated MasterProfile object.
"""

from pathlib import Path
from typing import Union, Dict, Any, Optional
import json
import logging

from models import MasterProfile
from llm_client import LLMClient

logger = logging.getLogger(__name__)


RESUME_BOOTSTRAP_SYSTEM_PROMPT = """
You are an expert resume parsing engine. You will be given the raw text of a candidate's resume.
Extract a complete, structured Master Profile matching the target schema.

Rules:
1. Extract all factual information faithfully:
   - Personal info (full_name, email, phone, location, linkedin_url, portfolio_url)
   - Professional summary / objective
   - Work experience: company, title, location, start_date (YYYY-MM or YYYY), end_date (YYYY-MM, YYYY, or null if current), is_current (bool)
   - For each bullet point in work experience: text, skills_tags (list of skills demonstrated), and metrics (list of quantified numbers/percentages)
   - Education: institution, degree, field, graduation_date (YYYY-MM or YYYY)
   - Skills: categorize into technical (languages, frameworks), tools (software, platforms, devops), soft (leadership, communication), languages (spoken)
   - Certifications: list of certification titles
   - Projects: any standalone projects listed on the resume
2. Assign clean IDs:
   - Work experiences: we_1, we_2, etc.
   - Bullets: b_1, b_2, etc.
   - Education: ed_1, ed_2, etc.
   - Projects: proj_1, proj_2, etc.
3. Do not invent or hallucinate information not present in the text.
4. Output valid JSON matching the MasterProfile schema.
"""


def extract_text_from_file(file_path: Union[str, Path]) -> str:
    """Extracts raw text content from PDF, DOCX, TXT, or MD files."""
    path = Path(file_path).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(str(path))
            pages_text = [page.extract_text() or "" for page in reader.pages]
            full_text = "\n\n".join(pages_text).strip()
            if not full_text:
                raise ValueError("PDF contains no selectable text (might be a scanned/image PDF).")
            return full_text
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from PDF: {e}")

    elif suffix == ".docx":
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        except Exception as e:
            raise RuntimeError(f"Failed to extract text from DOCX: {e}")

    elif suffix in [".txt", ".md", ".json"]:
        return path.read_text(encoding="utf-8", errors="ignore")

    else:
        # Attempt plain text read as fallback
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            raise ValueError(f"Unsupported file format: {suffix} ({e})")


def bootstrap_master_profile_from_text(raw_text: str, llm_client: Optional[LLMClient] = None) -> MasterProfile:
    """Uses LLM structured extraction to parse resume text into a MasterProfile instance."""
    client = llm_client or LLMClient()
    user_prompt = f"RESUME RAW TEXT:\n\n{raw_text}"

    extracted_dict = client.generate_json(
        system_prompt=RESUME_BOOTSTRAP_SYSTEM_PROMPT,
        user_prompt=user_prompt,
        temperature=0.0
    )

    profile = MasterProfile(**extracted_dict)
    profile.ensure_unique_ids()
    return profile


def bootstrap_master_profile_from_file(file_path: Union[str, Path], llm_client: Optional[LLMClient] = None) -> MasterProfile:
    """Extracts text from a resume file and bootstraps a validated MasterProfile."""
    text = extract_text_from_file(file_path)
    return bootstrap_master_profile_from_text(text, llm_client=llm_client)

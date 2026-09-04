"""
ATS Compliance Validator for Stage 7.
Performs deterministic structural and formatting checks on resume data, DOCX, and PDF files.
Encodes 2026 ATS guidelines (ATS_Guidelines_and_Prompt_Design.md §2).
"""

from pathlib import Path
from typing import Union, List, Optional
import re

from models import TailoredResume, ATSValidationResult, ATSValidationIssue

APPROVED_FONTS = {
    "arial", "calibri", "georgia", "garamond",
    "times new roman", "helvetica", "tahoma", "verdana"
}

STANDARD_SECTIONS = {"summary", "experience", "education", "skills", "certifications", "projects"}


def validate_tailored_resume_structure(tailored: TailoredResume) -> ATSValidationResult:
    """Validates structural ATS rules on the TailoredResume data object before document generation."""
    issues: List[ATSValidationIssue] = []
    score = 100

    # Rule 1: Dedicated Skills Section exists and has entries
    total_skills = (
        len(tailored.selected_skills.technical)
        + len(tailored.selected_skills.tools)
        + len(tailored.selected_skills.soft)
        + len(tailored.selected_skills.languages)
    )
    if total_skills == 0:
        issues.append(
            ATSValidationIssue(
                rule="Dedicated Skills Section",
                message="ATS parsers heavily weight dedicated skills sections; no skills found in selected_skills.",
                severity="warning"
            )
        )
        score -= 15

    # Rule 2: Work Experience entries and bullets
    if not tailored.selected_experience:
        issues.append(
            ATSValidationIssue(
                rule="Experience Section",
                message="No work experience entries found in tailored resume.",
                severity="error"
            )
        )
        score -= 30
    else:
        for exp in tailored.selected_experience:
            if not exp.bullets:
                issues.append(
                    ATSValidationIssue(
                        rule="Bullet Points",
                        message=f"Experience at '{exp.company}' has 0 bullet points.",
                        severity="warning"
                    )
                )
                score -= 5

    # Rule 3: Target job title is set
    if not tailored.target_job_title:
        issues.append(
            ATSValidationIssue(
                rule="Target Job Title",
                message="Target job title is missing; ATS parsers check for role alignment.",
                severity="warning"
            )
        )
        score -= 5

    passed = not any(i.severity == "error" for i in issues)
    return ATSValidationResult(passed=passed, score=max(0, score), issues=issues)


def validate_docx_file(docx_path: Union[str, Path]) -> ATSValidationResult:
    """Deterministic validation of generated DOCX file structure."""
    path = Path(docx_path)
    if not path.exists():
        return ATSValidationResult(
            passed=False,
            score=0,
            issues=[ATSValidationIssue(rule="File Exists", message=f"DOCX file not found: {path}", severity="error")]
        )

    try:
        import docx
        doc = docx.Document(str(path))
        issues: List[ATSValidationIssue] = []
        score = 100

        # Check 1: No tables used for overall layout
        if len(doc.tables) > 2:
            issues.append(
                ATSValidationIssue(
                    rule="No Layout Tables",
                    message="Detected multiple tables in document; tables for layout cause parsing scrambling in strict ATS systems.",
                    severity="warning"
                )
            )
            score -= 15

        # Check 2: Verify body text exists
        total_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if len(total_text) < 100:
            issues.append(
                ATSValidationIssue(
                    rule="Document Content",
                    message="Document text is unusually short (< 100 characters).",
                    severity="error"
                )
            )
            score -= 40

        passed = not any(i.severity == "error" for i in issues)
        return ATSValidationResult(passed=passed, score=max(0, score), issues=issues)

    except Exception as e:
        return ATSValidationResult(
            passed=False,
            score=0,
            issues=[ATSValidationIssue(rule="DOCX Parsing", message=f"Failed to inspect DOCX: {e}", severity="error")]
        )


def validate_pdf_file(pdf_path: Union[str, Path]) -> ATSValidationResult:
    """Validates that a PDF has an intact, selectable text layer (not a rasterized image)."""
    path = Path(pdf_path)
    if not path.exists():
        return ATSValidationResult(
            passed=False,
            score=0,
            issues=[ATSValidationIssue(rule="File Exists", message=f"PDF file not found: {path}", severity="error")]
        )

    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(str(path))
        issues: List[ATSValidationIssue] = []
        score = 100

        if len(reader.pages) == 0:
            return ATSValidationResult(
                passed=False,
                score=0,
                issues=[ATSValidationIssue(rule="PDF Pages", message="PDF has 0 pages.", severity="error")]
            )

        total_extracted_text = ""
        for page in reader.pages:
            total_extracted_text += (page.extract_text() or "") + " "

        total_words = len(total_extracted_text.split())
        if total_words < 50:
            issues.append(
                ATSValidationIssue(
                    rule="Selectable Text Layer",
                    message="PDF contains almost no selectable text words. This indicates an image/scanned PDF that ATS cannot parse!",
                    severity="error"
                )
            )
            score -= 70

        # Check reasonable page count (1 to 2 pages)
        if len(reader.pages) > 2:
            issues.append(
                ATSValidationIssue(
                    rule="Page Length",
                    message=f"Resume is {len(reader.pages)} pages; 1-2 pages recommended.",
                    severity="warning"
                )
            )
            score -= 10

        passed = not any(i.severity == "error" for i in issues)
        return ATSValidationResult(passed=passed, score=max(0, score), issues=issues)

    except Exception as e:
        return ATSValidationResult(
            passed=False,
            score=0,
            issues=[ATSValidationIssue(rule="PDF Parsing", message=f"Failed to inspect PDF: {e}", severity="error")]
        )
